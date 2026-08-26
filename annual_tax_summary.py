#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
annual_tax_summary.py — Annual tax filing summary.

Reads stock-sale capital-gains data ONLY from form1325_support_<year>.xlsx (Tax Data sheet).
Does NOT re-parse IBKR or EquatePlus source files for capital-gains data.
Reads IBKR PDF for dividends, interest, and withholding tax only.

Produces tax_annual_summary_<year>.xlsx with filing rows (1301/1322/1324)
and a Provenance sheet recording the SHA-256 hash of the consumed 1325 workbook.
"""

import argparse
import datetime
import glob
import hashlib
import os

import pandas as pd
from openpyxl.utils import get_column_letter

from tax_utils import _round_ils
from ibkr_parser import (
    extract_full_text,
    parse_ticker_tax_country,
    parse_dividends_by_country,
    parse_interest_transactions_ils,
    parse_withholding_tax_ils,
)

SCHEMA_VERSION = "1"
REQUIRED_COLUMNS = [
    "schema_version", "tax_year", "source", "fill_id", "lot_id",
    "gross_sale_proceeds_ils_exact",
    "taxable_gain_ils_filing", "deductible_loss_ils_filing",
    "tax_rate",
]


def load_tax_data(year: str) -> list[dict]:
    """
    Load and validate Tax Data from form1325_support_<year>.xlsx.
    Raises RuntimeError for schema mismatches or missing data.
    Raises SystemExit if the file is absent.
    """
    path = f"form1325_support_{year}.xlsx"
    if not os.path.exists(path):
        raise SystemExit(
            f"form1325_support_{year}.xlsx not found. "
            f"Run: python generate_1325_support.py --year {year}"
        )
    df = pd.read_excel(path, sheet_name="Tax Data")
    meta_raw = pd.read_excel(path, sheet_name="Metadata", header=None, index_col=0)
    # Normalize: pandas/openpyxl may read numeric-looking strings as int/float
    meta = {str(k).strip(): str(v).strip() for k, v in meta_raw[1].to_dict().items()}

    if meta.get("schema_version") != SCHEMA_VERSION:
        raise RuntimeError(
            f"Unsupported Tax Data schema version {meta.get('schema_version')!r}. "
            f"Re-run: python generate_1325_support.py --year {year}"
        )
    if meta.get("tax_year") != str(year).strip():
        raise RuntimeError(
            f"Tax Data tax_year {meta.get('tax_year')!r} != requested year {year!r}. "
            f"Re-run: python generate_1325_support.py --year {year}"
        )
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise RuntimeError(
            f"Tax Data missing columns: {missing}. "
            f"Re-run: python generate_1325_support.py --year {year}"
        )
    if df["lot_id"].duplicated().any():
        dupes = df.loc[df["lot_id"].duplicated(keep=False), "lot_id"].unique().tolist()
        raise RuntimeError(f"Duplicate lot_id values in Tax Data: {dupes}")

    return df.to_dict("records")


def _sha256_file(path: str) -> str:
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def _verify_ibkr_pdf_hash(ibkr_pdf: str, wbook_path: str) -> None:
    """
    Check that the IBKR PDF used for income parsing matches the one whose
    capital-gains lots were recorded in the 1325 workbook Sources sheet.
    Raises RuntimeError on hash mismatch; warns (does not fail) if Sources
    sheet is absent or has no IBKR row (e.g. EquatePlus-only year).
    """
    if not os.path.exists(wbook_path):
        return  # load_tax_data already checked this; unreachable in normal flow
    try:
        sources_df = pd.read_excel(wbook_path, sheet_name="Sources")
    except Exception:
        print(f"  WARNING: could not read Sources sheet from {wbook_path} — skipping hash check")
        return

    ibkr_rows = sources_df[sources_df["type"] == "IBKR"] if "type" in sources_df.columns else pd.DataFrame()
    if ibkr_rows.empty:
        print("  INFO: no IBKR entry in Sources sheet (EquatePlus-only workbook) — skipping hash check")
        return

    recorded_hash = str(ibkr_rows.iloc[0]["sha256"]).strip()
    actual_hash = _sha256_file(ibkr_pdf)
    if actual_hash != recorded_hash:
        raise RuntimeError(
            f"IBKR PDF hash mismatch!\n"
            f"  Supplied: {os.path.basename(ibkr_pdf)} — sha256={actual_hash[:16]}...\n"
            f"  Recorded in {wbook_path} Sources: {recorded_hash[:16]}...\n"
            f"The PDF used for income parsing differs from the one used to generate Tax Data. "
            f"Re-run: python generate_1325_support.py --year {os.path.basename(wbook_path).replace('form1325_support_','').replace('.xlsx','')} "
            f"--ibkr-pdf {ibkr_pdf}"
        )
    print(f"  IBKR PDF hash verified against Sources sheet ({actual_hash[:16]}...)")


def _autosize_columns(ws) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 8), 40)


def main():
    parser = argparse.ArgumentParser(description="Generate annual tax filing summary")
    parser.add_argument("--year", required=True, help="Tax year (YYYY)")
    parser.add_argument("--ibkr-pdf", help="Path to IBKR *_tax_statement.pdf (overrides glob)")
    args = parser.parse_args()
    year = args.year

    # Load stock-sale data from Tax Data (no Moses re-calculation)
    print(f"Loading Tax Data from form1325_support_{year}.xlsx ...")
    rows = load_tax_data(year)
    print(f"  {len(rows)} lot rows loaded")

    # Aggregate by (source, tax_rate) — produces distinct 1322 buckets per rate
    from collections import defaultdict
    by_src_rate: dict = defaultdict(lambda: {"gains": 0.0, "losses": 0.0, "gross_exact": 0.0})
    for r in rows:
        key = (r["source"], float(r["tax_rate"]))
        by_src_rate[key]["gains"]       += r["taxable_gain_ils_filing"]
        by_src_rate[key]["losses"]      += r["deductible_loss_ils_filing"]
        by_src_rate[key]["gross_exact"] += r["gross_sale_proceeds_ils_exact"]

    supported_rates = {0.25}
    all_rates = {float(r["tax_rate"]) for r in rows}
    unsupported = all_rates - supported_rates
    if unsupported:
        raise RuntimeError(
            f"Tax Data contains unsupported tax rate(s): {sorted(unsupported)}. "
            f"Only {sorted(supported_rates)} are handled by this script. "
            f"Re-run generate_1325_support.py and verify the source data."
        )

    def _src_rate(source: str, rate: float) -> dict:
        return by_src_rate.get((source, rate), {"gains": 0.0, "losses": 0.0, "gross_exact": 0.0})

    ibkr_25   = _src_rate("IBKR", 0.25)
    equate_25 = _src_rate("EquatePlus", 0.25)

    ibkr_gains        = int(ibkr_25["gains"])
    ibkr_losses       = int(ibkr_25["losses"])
    ibkr_gross        = _round_ils(ibkr_25["gross_exact"])
    equate_gains      = int(equate_25["gains"])
    equate_losses     = int(equate_25["losses"])
    equate_gross      = _round_ils(equate_25["gross_exact"])
    combined_gains    = ibkr_gains + equate_gains
    combined_losses   = ibkr_losses + equate_losses
    combined_gross    = _round_ils(sum(r["gross_sale_proceeds_ils_exact"] for r in rows))

    print(f"  IBKR: gains={ibkr_gains}  losses={ibkr_losses}  gross={ibkr_gross}")
    print(f"  EquatePlus: gains={equate_gains}  losses={equate_losses}  gross={equate_gross}")
    print(f"  Combined: gains={combined_gains}  losses={combined_losses}  gross={combined_gross}")

    # IBKR income parsing (dividends, interest, WHT)
    wbook_path = f"form1325_support_{year}.xlsx"
    ibkr_pdf = None
    if args.ibkr_pdf:
        if not os.path.exists(args.ibkr_pdf):
            raise SystemExit(f"--ibkr-pdf not found: {args.ibkr_pdf}")
        ibkr_pdf = args.ibkr_pdf
    else:
        pdfs = glob.glob("*_tax_statement.pdf")
        if len(pdfs) > 1:
            raise SystemExit(f"Expected exactly 1 *_tax_statement.pdf, found {len(pdfs)}: {pdfs}. Use --ibkr-pdf.")
        if pdfs:
            ibkr_pdf = pdfs[0]

    if ibkr_pdf:
        _verify_ibkr_pdf_hash(ibkr_pdf, wbook_path)

        print(f"\nParsing dividends from {ibkr_pdf} ...")
        ticker_tax_country = parse_ticker_tax_country(ibkr_pdf)
        dividends_by_country = parse_dividends_by_country(ibkr_pdf, ticker_tax_country)
        dividend_total = sum(dividends_by_country.values())

        print("Parsing interest income ...")
        interest_total = parse_interest_transactions_ils(ibkr_pdf)

        print("Parsing withholding tax ...")
        text = extract_full_text(ibkr_pdf)
        wht_ils = parse_withholding_tax_ils(text)

        print(f"  Dividends: { {k: _round_ils(v) for k, v in dividends_by_country.items()} }")
        print(f"  Dividend total: {_round_ils(dividend_total)}")
        print(f"  Interest total: {_round_ils(interest_total)}")
        print(f"  WHT: {wht_ils}")
    else:
        print("\nNo IBKR PDF found — dividends, interest, and withholding tax will be zero.")
        dividends_by_country = {}
        dividend_total = 0.0
        interest_total = 0.0
        wht_ils = 0

    div_plus_income = dividend_total + interest_total

    # Build summary rows
    summary_rows = [
        {"Label": "Interest income",                  "Value (ILS)": _round_ils(interest_total)},
        {"Label": "Dividend spread by country",       "Value (ILS)": ""},
    ]
    for country, amount in dividends_by_country.items():
        summary_rows.append({"Label": f"  {country}", "Value (ILS)": _round_ils(amount)})
    summary_rows += [
        {"Label": "Dividend total",                   "Value (ILS)": _round_ils(dividend_total)},
        {"Label": "Dividend + external income total", "Value (ILS)": _round_ils(div_plus_income)},
        {"Label": "Foreign withholding tax",          "Value (ILS)": wht_ils},
        {"Label": "",                                 "Value (ILS)": ""},
        {"Label": "IBKR taxable realized stock gains",       "Value (ILS)": ibkr_gains},
        {"Label": "IBKR deductible realized stock losses",   "Value (ILS)": ibkr_losses},
        {"Label": "IBKR gross sale value (stock disposals)", "Value (ILS)": ibkr_gross},
        {"Label": "",                                        "Value (ILS)": ""},
        {"Label": "EquatePlus taxable realized gains",       "Value (ILS)": equate_gains},
        {"Label": "EquatePlus deductible realized losses",   "Value (ILS)": equate_losses},
        {"Label": "EquatePlus gross sale value",             "Value (ILS)": equate_gross},
        {"Label": "",                                        "Value (ILS)": ""},
        {"Label": "TOTAL taxable gains (IBKR + EquatePlus)",     "Value (ILS)": combined_gains},
        {"Label": "TOTAL deductible losses (IBKR + EquatePlus)", "Value (ILS)": combined_losses},
        {"Label": "TOTAL gross sale value (IBKR + EquatePlus)",  "Value (ILS)": combined_gross},
    ]

    out = f"tax_annual_summary_{year}.xlsx"
    print(f"\nWriting {out} ...")

    wbook_hash = _sha256_file(wbook_path)

    df = pd.DataFrame(summary_rows)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Annual Summary")
        pd.DataFrame().to_excel(writer, sheet_name="Provenance", index=False)

        wb = writer.book

        ws_summary = wb["Annual Summary"]
        ws_summary.column_dimensions["A"].width = 45
        ws_summary.column_dimensions["B"].width = 18

        ws_prov = wb["Provenance"]
        ws_prov.delete_rows(1, ws_prov.max_row)
        provenance_pairs = [
            ("source_1325_workbook", wbook_path),
            ("sha256_1325_workbook", wbook_hash),
            ("ibkr_pdf",             os.path.basename(ibkr_pdf) if ibkr_pdf else "none"),
            ("generation_timestamp", datetime.datetime.now().isoformat(timespec="seconds")),
        ]
        for key, val in provenance_pairs:
            ws_prov.append([key, val])
        _autosize_columns(ws_prov)

    print(f"Done. Wrote {out}")
    print(f"  1325 workbook hash recorded in Provenance: {wbook_hash[:16]}...")
    if not ibkr_pdf:
        print(
            "\n  WARNING: No IBKR PDF was supplied or found. "
            "Interest income, dividends, and withholding tax are recorded as zero. "
            "If you have IBKR income, re-run with: --ibkr-pdf <path>"
        )


if __name__ == "__main__":
    main()
